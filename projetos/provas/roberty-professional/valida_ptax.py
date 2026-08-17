"""
Valida a extracao do PTAX exatamente como o robo da certificacao deve fazer:
mesma pagina, mesmos seletores do enunciado, e confere contra a API oficial do BCB.

Enunciado: cotacoes do dolar de 01/11/2024 a 30/11/2024, gravando Data, Compra e Venda.
Seletores dados na prova:
    linhas  -> tr[class*="fundoPadraoBClaro"]
    Data    -> td:nth-of-type(1)
    Compra  -> td:nth-of-type(3)
    Venda   -> td:nth-of-type(4)
"""
import shutil
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

AQUI = Path(__file__).parent
MODELO = Path(r"C:\Users\junio\Downloads\dev_digihub\material\EquipeTecnica\Equipe Técnica\Roberty\Certificação\output.xlsx")
SAIDA = AQUI / "output.xlsx"

DATA_INI = "01/11/2024"
DATA_FIM = "30/11/2024"
MOEDA_DOLAR = "61"

URL_CONSULTA = "https://ptax.bcb.gov.br/ptax_internet/consultaBoletim.do?method=consultarBoletim"
URL_OLINDA = (
    "https://olinda.bcb.gov.br/olinda/servico/PTAX/versao/v1/odata/"
    "CotacaoDolarPeriodo(dataInicial=@dataInicial,dataFinalCotacao=@dataFinalCotacao)"
)


def busca_no_site():
    """Uma unica requisicao, informando data inicial e final, como a prova exige."""
    payload = {
        "RadOpcao": "1",
        "DATAINI": DATA_INI,
        "DATAFIM": DATA_FIM,
        "ChkMoeda": MOEDA_DOLAR,
    }
    r = requests.post(URL_CONSULTA, data=payload, timeout=60)
    r.raise_for_status()
    r.encoding = r.apparent_encoding

    sopa = BeautifulSoup(r.text, "html.parser")
    linhas = sopa.select('tr[class*="fundoPadraoBClaro"]')

    cotacoes = []
    for linha in linhas:
        celulas = linha.find_all("td")
        if len(celulas) < 4:
            continue
        cotacoes.append({
            "data": celulas[0].get_text(strip=True),
            "compra": celulas[2].get_text(strip=True),
            "venda": celulas[3].get_text(strip=True),
        })
    return cotacoes


def busca_na_api():
    """Fonte independente para conferir os valores."""
    params = {
        "@dataInicial": f"'11-01-2024'",
        "@dataFinalCotacao": f"'11-30-2024'",
        "$format": "json",
        "$select": "cotacaoCompra,cotacaoVenda,dataHoraCotacao",
    }
    r = requests.get(URL_OLINDA, params=params, timeout=60)
    r.raise_for_status()
    saida = {}
    for item in r.json()["value"]:
        dia = item["dataHoraCotacao"][:10]
        dia_br = f"{dia[8:10]}/{dia[5:7]}/{dia[0:4]}"
        saida[dia_br] = (round(item["cotacaoCompra"], 4), round(item["cotacaoVenda"], 4))
    return saida


def ultima_linha_com_dado(ws):
    """O modelo vem com linhas vazias ja formatadas, entao max_row mente.
    'Preencher ultima linha' do Roberty grava apos a ultima linha COM conteudo."""
    ultima = 0
    for linha in range(1, ws.max_row + 1):
        if ws.cell(row=linha, column=1).value not in (None, ""):
            ultima = linha
    return ultima


def preenche_planilha(cotacoes):
    shutil.copy(MODELO, SAIDA)
    wb = load_workbook(SAIDA)
    ws = wb.active
    linha = ultima_linha_com_dado(ws) + 1
    for c in cotacoes:
        ws.cell(row=linha, column=1, value=c["data"])
        ws.cell(row=linha, column=2, value=c["compra"])
        ws.cell(row=linha, column=3, value=c["venda"])
        linha += 1
    wb.save(SAIDA)
    return linha - 1


def main():
    print(f"consultando PTAX de {DATA_INI} a {DATA_FIM}...", flush=True)
    site = busca_no_site()
    print(f"linhas capturadas pelos seletores da prova: {len(site)}")
    if not site:
        raise SystemExit("nenhuma linha capturada, seletor nao bateu")

    print(f"primeira: {site[0]}")
    print(f"ultima:   {site[-1]}")

    print("\nconferindo com a API oficial do BCB...")
    api = busca_na_api()
    print(f"dias na API: {len(api)}")

    erros = []
    for c in site:
        esperado = api.get(c["data"])
        if esperado is None:
            erros.append(f"{c['data']}: nao existe na API")
            continue
        compra = float(c["compra"].replace(",", "."))
        venda = float(c["venda"].replace(",", "."))
        if abs(compra - esperado[0]) > 1e-4 or abs(venda - esperado[1]) > 1e-4:
            erros.append(f"{c['data']}: site {compra}/{venda} != api {esperado[0]}/{esperado[1]}")

    if len(site) != len(api):
        erros.append(f"quantidade divergente: site {len(site)} x api {len(api)}")

    if erros:
        print("\nDIVERGENCIAS:")
        for e in erros:
            print("  " + e)
    else:
        print("OK: todas as linhas batem com a API oficial")

    ultima = preenche_planilha(site)
    print(f"\nplanilha gravada: {SAIDA}")
    print(f"dados nas linhas 2 a {ultima} ({ultima - 1} cotacoes)")


if __name__ == "__main__":
    main()
